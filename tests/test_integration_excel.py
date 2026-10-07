from __future__ import annotations

from collections import Counter
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from src.config import (
    DEFAULT_INPUT_XLSX,
    EXEC_DATA_QUALITY_REVIEW,
    EXEC_READY,
    STATUS_AUTO_DISCOUNT,
)
from src.loader import EXPECTED_HEADERS, load_items
from src.pipeline import run_pipeline

BUSINESS_DATE = date(2026, 10, 7)


def _copy_reordered_source(source: Path, target: Path) -> None:
    src_wb = load_workbook(source, data_only=True)
    src_ws = src_wb["Остатки_Кулинария_19_30"]
    headers = [c.value for c in src_ws[1]]
    reordered = list(reversed(EXPECTED_HEADERS))
    positions = {h: headers.index(h) + 1 for h in EXPECTED_HEADERS}

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки_Кулинария_19_30"
    ws.append(reordered)
    for row in src_ws.iter_rows(min_row=2, values_only=True):
        ws.append([row[positions[h] - 1] for h in reordered])
    wb.save(target)
    src_wb.close()


def test_real_export_business_counts_and_output(tmp_path: Path) -> None:
    output = tmp_path / "Итоги_Вечерней_Уценки_20_00.xlsx"
    result = run_pipeline(input_path=DEFAULT_INPUT_XLSX, output_path=output, business_date=BUSINESS_DATE)

    assert result.input_rows == 107
    assert result.kpi.processed_rows == 107
    assert result.kpi.auto_discount == 50
    assert result.kpi.block_category == 31
    assert result.kpi.block_expiry == 19
    assert result.kpi.block_brak == 1
    assert result.kpi.margin_alert == 6
    assert result.kpi.execution_ready == 43
    assert result.kpi.data_quality_review == 7

    business_counts = Counter(d.business_status for d in result.decisions)
    assert business_counts[STATUS_AUTO_DISCOUNT] == 50
    assert sum(v for k, v in business_counts.items() if k is not None) == 107
    assert sum(1 for d in result.decisions if d.business_status == STATUS_AUTO_DISCOUNT and d.execution_status == EXEC_READY) == 43
    assert sum(1 for d in result.decisions if d.business_status == STATUS_AUTO_DISCOUNT and d.execution_status == EXEC_DATA_QUALITY_REVIEW) == 7

    assert result.kpi.auto_stock_retail_value == 626065
    assert result.kpi.potential_revenue_after_discount == 437127.5
    assert result.kpi.auto_cost_value == 331406
    assert result.kpi.remainder_over_cost == 105721.5
    assert result.kpi.ready_stock_retail_value == 538730
    assert result.kpi.ready_revenue == 376221
    assert result.kpi.ready_cost_value == 286686
    assert result.kpi.ready_remainder_over_cost == 89535
    assert result.kpi.master_data_revenue_at_risk == 60906.5

    wb = load_workbook(output)
    assert wb.sheetnames == ["01_Директор", "02_План_уценки", "03_Исключения", "04_Контроль_данных"]
    plan = wb["02_План_уценки"]
    plan_rows = sum(1 for row in plan.iter_rows(min_row=2, max_col=1, values_only=True) if row[0])
    assert plan_rows == 43
    for row in plan.iter_rows(min_row=2, min_col=2, max_col=2):
        cell = row[0]
        if cell.value in (None, ""):
            continue
        assert cell.number_format == "@"
        assert not isinstance(cell.value, float)

    exceptions = wb["03_Исключения"]
    exc_rows = sum(1 for row in exceptions.iter_rows(min_row=2, max_col=1, values_only=True) if row[0])
    assert exc_rows == 64

    assert result.print_path.exists()
    print_wb = load_workbook(result.print_path)
    assert print_wb.sheetnames == ["Лист_для_печати"]
    print_ws = print_wb["Лист_для_печати"]
    card_titles = sum(
        1
        for row in print_ws.iter_rows(values_only=True)
        for value in row
        if value == "ВЕЧЕРНЯЯ СКИДКА −30%"
    )
    assert card_titles == 43
    printed_text = "\n".join(
        str(value)
        for row in print_ws.iter_rows(values_only=True)
        for value in row
        if value is not None
    )
    assert "Стейк из семги с овощами-гриль" not in printed_text  # MARGIN_ALERT, not READY
    assert "Салат Цезарь с курицей" in printed_text
    assert print_ws.page_setup.orientation == "landscape"


def test_reordered_columns_produce_same_business_counts(tmp_path: Path) -> None:
    reordered = tmp_path / "reordered.xlsx"
    output = tmp_path / "out.xlsx"
    _copy_reordered_source(DEFAULT_INPUT_XLSX, reordered)
    result = run_pipeline(input_path=reordered, output_path=output, business_date=BUSINESS_DATE)
    assert result.kpi.auto_discount == 50
    assert result.kpi.block_expiry == 19
    assert result.kpi.execution_ready == 43


def test_missing_required_column_fails_clearly(tmp_path: Path) -> None:
    path = tmp_path / "missing.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки_Кулинария_19_30"
    ws.append([h for h in EXPECTED_HEADERS if h != "Себестоимость_тг"])
    wb.save(path)
    with pytest.raises(ValueError, match="Себестоимость_тг"):
        load_items(path)
