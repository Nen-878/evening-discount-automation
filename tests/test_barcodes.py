from __future__ import annotations

from datetime import time
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.config import EXEC_DATA_QUALITY_REVIEW, EXEC_READY, STATUS_AUTO_DISCOUNT
from src.decision_engine import decide_item
from src.excel_report import write_plan_sheet
from src.loader import cell_as_text
from tests.factories import BUSINESS_DATE, make_item, with_duplicate_barcodes


def test_barcode_kept_as_string() -> None:
    item = make_item(barcode="4870001001015")
    assert isinstance(item.barcode, str)
    assert item.barcode == "4870001001015"
    assert cell_as_text("4870001001015") == "4870001001015"


def test_missing_barcode_keeps_business_auto_but_blocks_execution() -> None:
    item = make_item(
        barcode=None,
        produced_time=time(10, 0),
        shelf_hours=Decimal("24"),
        category="Кулинария: Горячие блюда",
    )
    decision = decide_item(item, BUSINESS_DATE)
    assert decision.business_status == STATUS_AUTO_DISCOUNT
    assert decision.execution_status == EXEC_DATA_QUALITY_REVIEW


def test_ambiguous_multi_barcode_blocks_execution_only() -> None:
    item = make_item(barcode="4870001001121;4870001001128")
    decision = decide_item(item, BUSINESS_DATE)
    assert decision.business_status == STATUS_AUTO_DISCOUNT
    assert decision.execution_status == EXEC_DATA_QUALITY_REVIEW


def test_scientific_barcode_is_not_reconstructed() -> None:
    item = make_item(barcode="4.60701E+12", category="Кулинария: Салаты незаправленные")
    decision = decide_item(item, BUSINESS_DATE)
    assert decision.business_status == STATUS_AUTO_DISCOUNT
    assert decision.execution_status == EXEC_DATA_QUALITY_REVIEW
    assert item.barcode == "4.60701E+12"


def test_whitespace_barcode_is_trimmed_safely_and_ready() -> None:
    item = make_item(barcode=" 4870001001060 ")
    decision = decide_item(item, BUSINESS_DATE)
    assert item.barcode == "4870001001060"
    assert item.raw.barcode_raw == " 4870001001060 "
    assert decision.business_status == STATUS_AUTO_DISCOUNT
    assert decision.execution_status == EXEC_READY
    assert any(issue.code == "BARCODE_WHITESPACE" for issue in item.issues)


def test_duplicate_barcode_blocks_execution_for_auto_goods() -> None:
    a = make_item(source_row=18, name="Салат А", barcode="4.60701E+12", code_1c="1")
    b = make_item(source_row=77, name="Шашлычки", barcode="4.60701E+12", code_1c="2", category="Кулинария: Горячие блюда")
    with_duplicate_barcodes([a, b])
    assert decide_item(a, BUSINESS_DATE).execution_status == EXEC_DATA_QUALITY_REVIEW
    assert decide_item(b, BUSINESS_DATE).execution_status == EXEC_DATA_QUALITY_REVIEW


def test_excel_writes_barcode_as_text(tmp_path: Path) -> None:
    item = make_item(
        barcode="4870001001015",
        produced_time=time(8, 30),
        qty=Decimal("5"),
        retail_price=Decimal("2200"),
        cost=Decimal("1200"),
    )
    decision = decide_item(item, BUSINESS_DATE)
    wb = Workbook()
    ws = wb.active
    write_plan_sheet(ws, [decision])
    path = tmp_path / "plan.xlsx"
    wb.save(path)
    loaded = load_workbook(path)
    cell = loaded.active["B2"]
    assert cell.value == "4870001001015"
    assert cell.number_format == "@"
    assert not isinstance(cell.value, float)
