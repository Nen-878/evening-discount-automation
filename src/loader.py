from __future__ import annotations

import logging
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterator, Optional

from openpyxl import load_workbook

from src.config import EXPECTED_SHEET_NAME
from src.models import ParsedItem, QualityIssue, RawItem

LOGGER = logging.getLogger(__name__)

EXPECTED_HEADERS = [
    "Код_1С",
    "Штрихкод",
    "Наименование_товара",
    "Товарная_категория",
    "Остаток_на_19_30",
    "Ед_изм",
    "Цена_розничная_тг",
    "Себестоимость_тг",
    "Время_изготовления",
    "Срок_годности_часов",
]


def cell_as_text(value: object) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, datetime):
        if value.hour == 0 and value.minute == 0 and value.second == 0 and value.microsecond == 0:
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, date) and not isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M:%S")
    if isinstance(value, float):
        as_int = int(value)
        if as_int == value:
            return str(as_int)
        return format(value, "g")
    text = str(value)
    return text if text != "" else None


def parse_decimal(value: object) -> tuple[Optional[Decimal], Optional[str]]:
    if value is None or value == "":
        return None, "пусто"
    if isinstance(value, bool):
        return None, "логическое значение вместо числа"
    if isinstance(value, Decimal):
        return value, None
    if isinstance(value, int):
        return Decimal(value), None
    if isinstance(value, float):
        try:
            return Decimal(str(value)), None
        except InvalidOperation:
            return None, f"некорректное число {value!r}"
    if isinstance(value, str):
        normalized = value.strip().replace(" ", "").replace(",", ".")
        if normalized == "":
            return None, "пусто"
        try:
            return Decimal(normalized), None
        except InvalidOperation:
            return None, f"некорректная строка числа {value!r}"
    try:
        return Decimal(str(value)), None
    except (InvalidOperation, ValueError, TypeError):
        return None, f"неподдерживаемый тип {type(value).__name__}"


def parse_production(value: object) -> tuple[Optional[time], Optional[date], Optional[datetime], Optional[str]]:
    if value is None or value == "":
        return None, None, None, "пусто"
    if isinstance(value, datetime):
        return value.time(), value.date(), value, None
    if isinstance(value, date) and not isinstance(value, datetime):
        return None, value, None, None
    if isinstance(value, time):
        return value, None, None, None
    if isinstance(value, str):
        text = value.strip()
        for fmt in ("%H:%M:%S", "%H:%M"):
            try:
                return datetime.strptime(text, fmt).time(), None, None, None
            except ValueError:
                continue
        for fmt in ("%Y-%m-%d", "%d.%m.%Y"):
            try:
                return None, datetime.strptime(text, fmt).date(), None, None
            except ValueError:
                continue
        return None, None, None, f"неразборчивое значение {value!r}"
    return None, None, None, f"неподдерживаемый тип {type(value).__name__}"


def _header_index(header_values: tuple[object, ...]) -> dict[str, int]:
    headers = [str(v).strip() if v is not None else "" for v in header_values]
    duplicates = {h for h in headers if h and headers.count(h) > 1}
    if duplicates:
        raise ValueError(f"В выгрузке дублируются заголовки: {sorted(duplicates)}")
    index = {header: i for i, header in enumerate(headers) if header}
    missing = [header for header in EXPECTED_HEADERS if header not in index]
    if missing:
        raise ValueError(
            "Структура выгрузки 1С изменилась. Не найдены обязательные колонки: "
            + ", ".join(missing)
        )
    return index


def iter_raw_rows(path: Path) -> Iterator[RawItem]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        if EXPECTED_SHEET_NAME not in workbook.sheetnames:
            raise ValueError(
                f"Не найден ожидаемый лист {EXPECTED_SHEET_NAME!r}. "
                f"Доступные листы: {', '.join(workbook.sheetnames)}"
            )
        sheet = workbook[EXPECTED_SHEET_NAME]
        rows = sheet.iter_rows(values_only=True)
        try:
            header_row = next(rows)
        except StopIteration as exc:
            raise ValueError("Выгрузка 1С пуста") from exc
        index = _header_index(header_row)

        def value(row: tuple[object, ...], header: str) -> object:
            pos = index[header]
            return row[pos] if pos < len(row) else None

        for excel_row_number, row in enumerate(rows, start=2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            yield RawItem(
                source_row=excel_row_number,
                code_1c_raw=value(row, "Код_1С"),
                barcode_raw=value(row, "Штрихкод"),
                name_raw=value(row, "Наименование_товара"),
                category_raw=value(row, "Товарная_категория"),
                qty_raw=value(row, "Остаток_на_19_30"),
                unit_raw=value(row, "Ед_изм"),
                retail_raw=value(row, "Цена_розничная_тг"),
                cost_raw=value(row, "Себестоимость_тг"),
                produced_raw=value(row, "Время_изготовления"),
                shelf_hours_raw=value(row, "Срок_годности_часов"),
            )
    finally:
        workbook.close()


def load_items(path: Path) -> list[ParsedItem]:
    items: list[ParsedItem] = []
    for raw in iter_raw_rows(path):
        try:
            items.append(parse_raw_item(raw))
        except Exception:
            LOGGER.exception("Строка Excel %s не разобрана целиком", raw.source_row)
            items.append(
                ParsedItem(
                    raw=raw,
                    code_1c=cell_as_text(raw.code_1c_raw) or "",
                    barcode=(str(raw.barcode_raw).strip() if raw.barcode_raw is not None else None),
                    name=cell_as_text(raw.name_raw) or "",
                    category=cell_as_text(raw.category_raw),
                    qty=None,
                    unit=cell_as_text(raw.unit_raw),
                    retail_price=None,
                    cost=None,
                    produced_time=None,
                    produced_date=None,
                    produced_datetime=None,
                    shelf_hours=None,
                    issues=[QualityIssue("ROW_PARSE_FAILURE", "Строка не разобрана; исходные значения сохранены")],
                )
            )
    return items


def parse_raw_item(raw: RawItem) -> ParsedItem:
    issues: list[QualityIssue] = []
    code_1c = cell_as_text(raw.code_1c_raw) or ""
    if not code_1c:
        issues.append(QualityIssue("EMPTY_CODE", "Пустой Код_1С"))

    raw_barcode_text = None if raw.barcode_raw is None else str(raw.barcode_raw)
    barcode = raw_barcode_text.strip() if raw_barcode_text is not None else None
    if isinstance(raw.barcode_raw, float):
        issues.append(
            QualityIssue(
                "BARCODE_NUMERIC_FLOAT",
                "Штрихкод прочитан как число с плавающей точкой; исходный полный код не восстанавливается",
            )
        )

    name = cell_as_text(raw.name_raw) or ""
    if not name:
        issues.append(QualityIssue("EMPTY_NAME", "Пустое наименование"))

    category = cell_as_text(raw.category_raw)
    if category is None or category.strip() == "":
        category = None
        issues.append(QualityIssue("EMPTY_CATEGORY", "Пустая товарная категория"))

    unit = cell_as_text(raw.unit_raw)
    if unit is None or unit.strip() == "":
        unit = None
        issues.append(QualityIssue("EMPTY_UNIT", "Пустая единица измерения"))

    qty, qty_err = parse_decimal(raw.qty_raw)
    if qty_err:
        issues.append(QualityIssue("QTY_UNPARSEABLE", f"Остаток: {qty_err}"))
    retail, retail_err = parse_decimal(raw.retail_raw)
    if retail_err:
        issues.append(QualityIssue("RETAIL_UNPARSEABLE", f"Розничная цена: {retail_err}"))
    cost, cost_err = parse_decimal(raw.cost_raw)
    if cost_err:
        issues.append(QualityIssue("COST_UNPARSEABLE", f"Себестоимость: {cost_err}"))
    produced_time, produced_date, produced_datetime, prod_err = parse_production(raw.produced_raw)
    if prod_err:
        issues.append(QualityIssue("PRODUCTION_UNPARSEABLE", f"Время изготовления: {prod_err}"))
    shelf_hours, hours_err = parse_decimal(raw.shelf_hours_raw)
    if hours_err:
        issues.append(QualityIssue("SHELF_HOURS_UNPARSEABLE", f"Срок годности (часов): {hours_err}"))

    return ParsedItem(
        raw=raw,
        code_1c=code_1c,
        barcode=barcode,
        name=name,
        category=category,
        qty=qty,
        unit=unit,
        retail_price=retail,
        cost=cost,
        produced_time=produced_time,
        produced_date=produced_date,
        produced_datetime=produced_datetime,
        shelf_hours=shelf_hours,
        issues=issues,
    )
