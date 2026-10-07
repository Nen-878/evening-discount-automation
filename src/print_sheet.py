from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

from src.config import EXEC_READY, STATUS_AUTO_DISCOUNT
from src.models import Decision

DARK = "172033"
BLUE = "1D4ED8"
YELLOW = "FFF2CC"
PALE_YELLOW = "FFF9E6"
BORDER_COLOR = "B8C0CC"
WHITE = "FFFFFF"

THIN = Border(
    left=Side(style="thin", color=BORDER_COLOR),
    right=Side(style="thin", color=BORDER_COLOR),
    top=Side(style="thin", color=BORDER_COLOR),
    bottom=Side(style="thin", color=BORDER_COLOR),
)


def _fmt_number(value: Decimal | None, decimals: int = 1) -> str:
    if value is None:
        return "—"
    pattern = f"{{:,.{decimals}f}}"
    return pattern.format(float(value)).replace(",", " ").replace(".", ",")


def _fmt_money(value: Decimal | None) -> str:
    if value is None:
        return "—"
    # Promo prices are rounded to whole tenge / step 5, so no kopecks are useful on print cards.
    return f"{int(value):,}".replace(",", " ") + " ₸"


def _fmt_qty(value: Decimal | None, unit: str | None) -> str:
    if value is None:
        return "—"
    raw = _fmt_number(value, 3).rstrip("0").rstrip(",")
    return f"{raw} {unit or ''}".strip()


def _ready(decision: Decision) -> bool:
    return decision.business_status == STATUS_AUTO_DISCOUNT and decision.execution_status == EXEC_READY


def _style_card(ws, start_row: int, start_col: int, decision: Decision) -> None:
    end_col = start_col + 3
    item = decision.item

    # Card title
    ws.merge_cells(start_row=start_row, start_column=start_col, end_row=start_row, end_column=end_col)
    title = ws.cell(start_row, start_col, "ВЕЧЕРНЯЯ СКИДКА −30%")
    title.font = Font(name="Calibri", size=14, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=BLUE)
    title.alignment = Alignment(horizontal="center", vertical="center")

    # Product name across two visual rows
    ws.merge_cells(start_row=start_row + 1, start_column=start_col, end_row=start_row + 2, end_column=end_col)
    name = ws.cell(start_row + 1, start_col, item.name)
    name.font = Font(name="Calibri", size=13, bold=True, color=DARK)
    name.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # Barcode number as text. We intentionally do NOT generate a barcode graphic because symbology is not specified.
    ws.merge_cells(start_row=start_row + 3, start_column=start_col, end_row=start_row + 3, end_column=end_col)
    barcode = ws.cell(start_row + 3, start_col, f"Штрихкод из 1С (номер): {item.barcode_display}")
    barcode.font = Font(name="Consolas", size=11, bold=True, color=DARK)
    barcode.alignment = Alignment(horizontal="center", vertical="center")
    barcode.number_format = "@"

    # Old/new prices
    ws.merge_cells(start_row=start_row + 4, start_column=start_col, end_row=start_row + 4, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 4, start_column=start_col + 2, end_row=start_row + 4, end_column=end_col)
    old_cell = ws.cell(start_row + 4, start_col, f"Было: {_fmt_money(item.retail_price)}")
    old_cell.font = Font(name="Calibri", size=11, color="6B7280")
    old_cell.alignment = Alignment(horizontal="center", vertical="center")
    new_cell = ws.cell(start_row + 4, start_col + 2, f"НОВАЯ ЦЕНА: {_fmt_money(decision.discount_price)}")
    new_cell.font = Font(name="Calibri", size=16, bold=True, color=DARK)
    new_cell.fill = PatternFill("solid", fgColor=YELLOW)
    new_cell.alignment = Alignment(horizontal="center", vertical="center")

    # Qty / expiry
    ws.merge_cells(start_row=start_row + 5, start_column=start_col, end_row=start_row + 5, end_column=start_col + 1)
    ws.merge_cells(start_row=start_row + 5, start_column=start_col + 2, end_row=start_row + 5, end_column=end_col)
    qty_cell = ws.cell(start_row + 5, start_col, f"Остаток: {_fmt_qty(item.qty, item.unit)}")
    qty_cell.font = Font(name="Calibri", size=10, bold=True, color=DARK)
    qty_cell.alignment = Alignment(horizontal="center", vertical="center")
    expiry_text = decision.expiry_at.strftime("%d.%m.%Y %H:%M") if decision.expiry_at else "—"
    exp_cell = ws.cell(start_row + 5, start_col + 2, f"Продать до: {expiry_text}")
    exp_cell.font = Font(name="Calibri", size=10, bold=True, color=DARK)
    exp_cell.alignment = Alignment(horizontal="center", vertical="center")
    if decision.risk_flag:
        exp_cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)

    # Traceability
    ws.merge_cells(start_row=start_row + 6, start_column=start_col, end_row=start_row + 6, end_column=end_col)
    trace = ws.cell(start_row + 6, start_col, f"Код 1С: {item.code_1c}   •   1 карточка = 1 товарная позиция")
    trace.font = Font(name="Calibri", size=8, color="6B7280")
    trace.alignment = Alignment(horizontal="center", vertical="center")

    for r in range(start_row, start_row + 7):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(r, c)
            cell.border = THIN
            if cell.fill.fill_type is None:
                cell.fill = PatternFill("solid", fgColor=WHITE)

    for r, height in {
        start_row: 24,
        start_row + 1: 28,
        start_row + 2: 28,
        start_row + 3: 24,
        start_row + 4: 34,
        start_row + 5: 26,
        start_row + 6: 18,
    }.items():
        ws.row_dimensions[r].height = height


def build_print_workbook(decisions: list[Decision], output_path: Path) -> Path:
    ready = [d for d in decisions if _ready(d)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Лист_для_печати"
    ws.sheet_view.showGridLines = False

    # A:I: two 4-column cards separated by E.
    for col in ("A", "B", "C", "D", "F", "G", "H", "I"):
        ws.column_dimensions[col].width = 16
    ws.column_dimensions["E"].width = 3

    ws.merge_cells("A1:I1")
    ws["A1"] = "Лист для маркировки вечерней уценки"
    ws["A1"].font = Font(name="Calibri", size=16, bold=True, color=DARK)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:I2")
    ws["A2"] = (
        "Только позиции, готовые к уценке. Штрихкод показан номером из 1С; новый графический barcode не генерируется. "
        "Количество физических стикеров определяется сотрудником по фактической фасовке/контейнерам."
    )
    ws["A2"].font = Font(name="Calibri", size=9, color="6B7280")
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[2].height = 34

    base_row = 4
    cards_per_page = 4  # 2 columns × 2 rows
    card_height_with_gap = 8

    for idx, decision in enumerate(ready):
        within_page = idx % cards_per_page
        page = idx // cards_per_page
        page_start = base_row + page * (2 * card_height_with_gap)
        block_row = page_start + (within_page // 2) * card_height_with_gap
        block_col = 1 if within_page % 2 == 0 else 6
        _style_card(ws, block_row, block_col, decision)

    total_pages = (len(ready) + cards_per_page - 1) // cards_per_page if ready else 1
    for page in range(1, total_pages):
        break_row = base_row + page * (2 * card_height_with_gap)
        ws.row_breaks.append(Break(id=break_row))

    last_row = base_row + max(total_pages, 1) * (2 * card_height_with_gap) - 1
    ws.print_area = f"A1:I{last_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = "Страница &P из &N"
    ws.oddFooter.right.text = f"Готовых позиций: {len(ready)}"

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return output_path
